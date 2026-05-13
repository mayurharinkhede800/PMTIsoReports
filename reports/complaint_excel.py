from openpyxl import load_workbook
from datetime import datetime
import os

# ---------------------------------------------------------
# GLOBAL HELPER: Safe Replace (Ignores Read-Only Merged Cells)
# ---------------------------------------------------------
def safe_replace_all(sheet, key, value):
    value = str(value) if value is not None else ""
    for row in sheet.iter_rows():
        for cell in row:
            # Skip secondary merged cells to prevent Read-Only crashes
            if type(cell).__name__ == 'MergedCell':
                continue
            
            if cell.value and key in str(cell.value):
                cell.value = str(cell.value).replace(key, value)

def parse_time(value):
    if not value: return None
    try: return datetime.fromisoformat(str(value))
    except: return None

def f_date(dt): return dt.strftime("%d-%m-%Y") if dt else ""
def f_time(dt): return dt.strftime("%H:%M") if dt else ""

# ---------------------------------------------------------
# 1. SINGLE TICKET EXPORT
# ---------------------------------------------------------
def generate_complaint_excel(ticket, download_path):
    BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    TEMPLATE_PATH = os.path.join(BASE_DIR, "templates", "complaint_register.xlsx")

    if not os.path.exists(TEMPLATE_PATH):
        raise FileNotFoundError(f"Template not found at: {TEMPLATE_PATH}")

    wb = load_workbook(TEMPLATE_PATH)
    ws = wb.active

    break_dt = parse_time(ticket.get("ticket_raised_time"))
    end_dt = parse_time(ticket.get("ticket_completion_time"))

    bd_date, bd_time = f_date(break_dt), f_time(break_dt)
    rc_date, rc_time = f_date(end_dt), f_time(end_dt)

    safe_replace_all(ws, "{{ticket_no}}", ticket.get("ticket_no", ""))
    safe_replace_all(ws, "{{description}}", ticket.get("title", ""))
    safe_replace_all(ws, "{{machine_name}}", ticket.get("machine_name") or "")
    safe_replace_all(ws, "{{area}}", ticket.get("area") or "")
    safe_replace_all(ws, "{{raised_by}}", ticket.get("raised_by") or "")
    safe_replace_all(ws, "{{assigned_to}}", ticket.get("assigned_to") or "")
    safe_replace_all(ws, "{{bd_date}}", bd_date)
    safe_replace_all(ws, "{{bd_time}}", bd_time)
    safe_replace_all(ws, "{{rc_date}}", rc_date)
    safe_replace_all(ws, "{{rc_time}}", rc_time)
    safe_replace_all(ws, "{{action}}", "Maintenance Done")

    os.makedirs(download_path, exist_ok=True)
    file_path = os.path.join(download_path, f"complaint_{ticket.get('ticket_no')}.xlsx")
    wb.save(file_path)
    print(f"✅ Single Complaint Register Generated: {file_path}")

# ---------------------------------------------------------
# 2. BULK DATE RANGE EXPORT (MULTI-SHEET)
# ---------------------------------------------------------
def generate_complaint_excel_bulk(tickets, download_path, start_date, end_date):
    BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    TEMPLATE_PATH = os.path.join(BASE_DIR, "templates", "complaint_register.xlsx")

    if not os.path.exists(TEMPLATE_PATH):
        raise FileNotFoundError(f"Template not found at: {TEMPLATE_PATH}")

    wb = load_workbook(TEMPLATE_PATH)
    template_sheet = wb.active
    template_sheet.title = "Master_Template"

    for i, ticket in enumerate(tickets):
        new_sheet = wb.copy_worksheet(template_sheet)
        
        # Ensure title is safe and fits Excel's 31 char limit
        raw_title = str(ticket.get("ticket_no", f"Ticket_{i+1}"))
        safe_title = raw_title.replace("/", "-").replace("\\", "-").replace("?", "").replace("*", "")
        new_sheet.title = safe_title[:31]

        break_dt = parse_time(ticket.get("ticket_raised_time"))
        end_dt = parse_time(ticket.get("ticket_completion_time"))
        bd_date, bd_time = f_date(break_dt), f_time(break_dt)
        rc_date, rc_time = f_date(end_dt), f_time(end_dt)

        safe_replace_all(new_sheet, "{{ticket_no}}", ticket.get("ticket_no", ""))
        safe_replace_all(new_sheet, "{{description}}", ticket.get("title", ""))
        safe_replace_all(new_sheet, "{{machine_name}}", ticket.get("machine_name") or "")
        safe_replace_all(new_sheet, "{{area}}", ticket.get("area") or "")
        safe_replace_all(new_sheet, "{{raised_by}}", ticket.get("raised_by") or "")
        safe_replace_all(new_sheet, "{{assigned_to}}", ticket.get("assigned_to") or "")
        safe_replace_all(new_sheet, "{{bd_date}}", bd_date)
        safe_replace_all(new_sheet, "{{bd_time}}", bd_time)
        safe_replace_all(new_sheet, "{{rc_date}}", rc_date)
        safe_replace_all(new_sheet, "{{rc_time}}", rc_time)
        safe_replace_all(new_sheet, "{{action}}", "Maintenance Done")

    # Remove the blank template so the user only sees the filled tickets
    wb.remove(template_sheet)

    os.makedirs(download_path, exist_ok=True)
    safe_start = str(start_date).replace("/", "-")
    safe_end = str(end_date).replace("/", "-")
    
    file_path = os.path.join(download_path, f"Complaint_Register_Tabs_{safe_start}_to_{safe_end}.xlsx")
    wb.save(file_path)
    print(f"✅ Multi-Sheet Complaint Register Generated: {file_path}")