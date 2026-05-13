from docx import Document
from copy import deepcopy
from collections import defaultdict
from datetime import datetime
import os


# ---------------------------------------------------------
# CLEAR CELL COMPLETELY
# ---------------------------------------------------------
def clear_cell(cell):

    cell.text = ""

    for paragraph in cell.paragraphs:

        for run in paragraph.runs:

            run.text = ""


# ---------------------------------------------------------
# WRITE CELL
# ---------------------------------------------------------
def write_cell(cell, value):

    clear_cell(cell)

    cell.text = str(value) if value else ""


# ---------------------------------------------------------
# REPLACE TEXT
# ---------------------------------------------------------
def replace_text(cell, old, new):

    if old in cell.text:

        cell.text = cell.text.replace(
            old,
            str(new)
        )


# ---------------------------------------------------------
# GENERATE TOOLS & TACKLES REPORT
# ---------------------------------------------------------
def generate_tools_tackles_report(
    data,
    download_path,
    filename="Tools_And_Tackles_Report.docx"
):

    # ---------------------------------------------------------
    # CHECK DATA
    # ---------------------------------------------------------
    if not data:
        raise Exception("No tools data found")

    # ---------------------------------------------------------
    # TEMPLATE PATH
    # ---------------------------------------------------------
    BASE_DIR = os.path.dirname(
        os.path.dirname(os.path.abspath(__file__))
    )

    TEMPLATE_PATH = os.path.join(
        BASE_DIR,
        "templates",
        "ToolsAndTackles.docx"
    )

    if not os.path.exists(TEMPLATE_PATH):

        raise FileNotFoundError(
            f"Template not found: {TEMPLATE_PATH}"
        )

    # ---------------------------------------------------------
    # GROUP DATEWISE
    # ---------------------------------------------------------
    grouped = defaultdict(list)

    for item in data:

        grouped[item["taken_date"]].append(item)

    sorted_dates = sorted(
        grouped.keys(),
        key=lambda x: datetime.strptime(
            x,
            "%d-%m-%Y"
        )
    )

    # ---------------------------------------------------------
    # FINAL DOCUMENT
    # ---------------------------------------------------------
    final_doc = None
    first_page = True

    # ---------------------------------------------------------
    # EACH DATE = ONE PAGE
    # ---------------------------------------------------------
    for taken_date in sorted_dates:

        items = grouped[taken_date]

        # ---------------------------------------------------------
        # LOAD TEMPLATE
        # ---------------------------------------------------------
        doc = Document(TEMPLATE_PATH)

        # ---------------------------------------------------------
        # CHECK TABLE
        # ---------------------------------------------------------
        if len(doc.tables) == 0:
            raise Exception("No table found in template")

        table = doc.tables[0]

        # ---------------------------------------------------------
        # REPLACE DATE PLACEHOLDER
        # ---------------------------------------------------------
        for row in table.rows:

            for cell in row.cells:

                replace_text(
                    cell,
                    "{{taken_date}}",
                    taken_date
                )

        # ---------------------------------------------------------
        # START ROW
        # ---------------------------------------------------------
        # Placeholder row index
        start_row = 2

        # ---------------------------------------------------------
        # AVAILABLE ROWS
        # ---------------------------------------------------------
        max_rows = len(table.rows) - start_row

        items = items[:max_rows]

        # ---------------------------------------------------------
        # INSERT DATA
        # ---------------------------------------------------------
        for i, item in enumerate(items):

            row = table.rows[start_row + i]

            cells = row.cells

            # S.NO
            if len(cells) > 0:
                write_cell(
                    cells[0],
                    i + 1
                )

            # TOOL NAME
            if len(cells) > 1:
                write_cell(
                    cells[1],
                    item.get("tool_name", "")
                )

            # EMPLOYEE NAME
            if len(cells) > 2:
                write_cell(
                    cells[2],
                    item.get("employee_name", "")
                )

            # TAKEN TIME
            if len(cells) > 3:
                write_cell(
                    cells[3],
                    item.get("taken_hour", "")
                )

            # RETURN TIME
            if len(cells) > 4:
                write_cell(
                    cells[4],
                    item.get("return_hour", "")
                )

        # ---------------------------------------------------------
        # CLEAR UNUSED ROWS
        # ---------------------------------------------------------
        for row_idx in range(
            start_row + len(items),
            len(table.rows)
        ):

            row = table.rows[row_idx]

            for cell in row.cells:

                clear_cell(cell)

        # ---------------------------------------------------------
        # MERGE INTO FINAL DOC
        # ---------------------------------------------------------
        if first_page:

            final_doc = doc
            first_page = False

        else:

            final_doc.add_page_break()

            for element in doc.element.body:

                final_doc.element.body.append(
                    deepcopy(element)
                )

    # ---------------------------------------------------------
    # CREATE DOWNLOAD FOLDER
    # ---------------------------------------------------------
    os.makedirs(download_path, exist_ok=True)

    # ---------------------------------------------------------
    # SAVE FILE
    # ---------------------------------------------------------
    file_path = os.path.join(
        download_path,
        filename
    )

    final_doc.save(file_path)

    print("\n✅ Tools & Tackles Report Generated")
    print(f"📁 Saved at: {file_path}")

    return file_path