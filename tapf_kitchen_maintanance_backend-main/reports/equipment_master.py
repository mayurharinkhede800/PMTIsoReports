from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.properties import PageSetupProperties
import os
from copy import copy
from math import ceil

def generate_equipment_master(data, download_path, filename="Equipment_Master.xlsx"):
    if not data:
        print("❌ No equipment found")
        return

    BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    TEMPLATE_PATH = os.path.join(BASE_DIR, "templates", "equipment_master_template.xlsx")
    LOGO_PATH = os.path.join(BASE_DIR, "templates", "logo.png")

    ITEMS_PER_SHEET = 26
    start_row = 8

    total_sheets = ceil(len(data) / ITEMS_PER_SHEET)

    final_wb = load_workbook(TEMPLATE_PATH)
    final_wb.remove(final_wb.active)

    # -------------------------------
    # CLEAN + LIMIT TEXT
    # -------------------------------
    def clean_text(value, max_len=35):
        if not value:
            return ""
        text = str(value).replace("\n", " ").strip()
        return text[:max_len]

    # -------------------------------
    # SAFE WRITE (MERGED CELL SAFE)
    # -------------------------------
    def safe_write(ws, row, col, value, align="left"):
        cell = ws.cell(row=row, column=col)
        for merged in ws.merged_cells.ranges:
            if cell.coordinate in merged:
                cell = ws.cell(merged.min_row, merged.min_col)
                break

        cell.value = value
        cell.font = Font(size=11)
        cell.alignment = Alignment(
            horizontal="center" if align == "center" else "left",
            vertical="center",
            wrap_text=False
        )

    # -------------------------------
    # ADD LOGO
    # -------------------------------
    def add_logo(ws):
        if os.path.exists(LOGO_PATH):
            img = Image(LOGO_PATH)
            img.width = 170
            img.height = 85
            img.anchor = "A2"
            ws.add_image(img)

    # -------------------------------
    # GENERATE SHEETS
    # -------------------------------
    for sheet_idx in range(total_sheets):
        temp_wb = load_workbook(TEMPLATE_PATH)
        temp_ws = temp_wb.active

        ws = final_wb.create_sheet(title=f"Page_{sheet_idx + 1}")

        # PRINT SETTINGS
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        ws.page_setup.fitToHeight = 1
        ws.page_setup.fitToWidth = 1
        ws.print_options.horizontalCentered = True
        ws.print_title_rows = '1:7'
        ws.print_area = "A1:G36"

        # COPY TEMPLATE EXACTLY
        for row in temp_ws.iter_rows():
            for cell in row:
                new_cell = ws.cell(row=cell.row, column=cell.column)
                new_cell.value = cell.value

                if cell.has_style:
                    new_cell.font = copy(cell.font)
                    new_cell.border = copy(cell.border)
                    new_cell.fill = copy(cell.fill)
                    new_cell.number_format = copy(cell.number_format)
                    new_cell.alignment = copy(cell.alignment)

        for col in temp_ws.column_dimensions:
            ws.column_dimensions[col].width = temp_ws.column_dimensions[col].width
            
        ws.column_dimensions['B'].width = 18   
        ws.column_dimensions['C'].width = 25   

        for merged in temp_ws.merged_cells.ranges:
            ws.merge_cells(str(merged))

        add_logo(ws)

        # INSERT DATA
        start_i = sheet_idx * ITEMS_PER_SHEET
        end_i = min(start_i + ITEMS_PER_SHEET, len(data))
        page_data = data[start_i:end_i]

        for i, item in enumerate(page_data):
            row = start_row + i
            sr_no = start_i + i + 1

            safe_write(ws, row, 1, sr_no, "center")
            safe_write(ws, row, 2, clean_text(item.get("machine_name"), 30))
            safe_write(ws, row, 4, clean_text(item.get("model"), 25))
            safe_write(ws, row, 5, item.get("date_of_commision"), "center")
            safe_write(ws, row, 6, clean_text(item.get("area"), 25))
            safe_write(ws, row, 7, clean_text(item.get("remarks"), 20))

        for col in range(1, 8):
            ws.cell(row=7, column=col).font = Font(size=12, bold=True)

    os.makedirs(download_path, exist_ok=True)
    file_path = os.path.join(download_path, filename)
    final_wb.save(file_path)

    print(f"✅ Equipment Master Generated: {file_path}")